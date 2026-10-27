"""
app.py — Web dashboard entry point (upload-and-analyze mode).

Videos are uploaded by an authenticated user, analyzed in a background
thread (see video_processor.VideoJob), and the results are served to the
dashboard UI as JSON. There is no live camera stream in this version — see
README.md for the previous live-CCTV mode if you need it back.

Run:
    python app.py
Then open:
    http://localhost:5000   (demo login: admin / admin123)
"""
import csv
import io
import json
import os
from datetime import timedelta
from datetime import datetime
import random
import smtplib
import socket
from email.mime.text import MIMEText
import qrcode

from flask import (
    Flask, render_template, jsonify, request, redirect,
    url_for, session, send_from_directory, send_file,
)
from datetime import datetime
from werkzeug.utils import secure_filename
from ultralytics import YOLO

import config
from auth import verify_credentials, login_required, current_user, current_user_id
from utils.alert_logger import init_log_file
from video_processor import VideoJob, new_job_id
from models import (
    db,
    User,
    EmergencyContact,
    AlertSetting,
    AlertHistory,
    UploadedVideo,
)
from utils import alerts

app = Flask(__name__)
otp_storage = {}

app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///crowd_alert.db"
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False



with app.app_context():
    db.init_app(app)
    db.create_all()
app.secret_key = config.SECRET_KEY
app.permanent_session_lifetime = timedelta(hours=config.SESSION_LIFETIME_HOURS)
app.config["MAX_CONTENT_LENGTH"] = config.MAX_UPLOAD_MB * 1024 * 1024

for _d in (config.UPLOAD_DIR, config.PROCESSED_DIR, config.KEYFRAME_DIR,
           config.HEATMAP_DIR, config.ANALYSIS_DIR, config.REPORT_DIR):
    os.makedirs(_d, exist_ok=True)

print(f"[INFO] Loading YOLO model: {config.MODEL_PATH} ...")
_shared_model = YOLO(config.MODEL_PATH)
_shared_fallback_model = YOLO(getattr(config, 'FALLBACK_MODEL_PATH', 'yolov8n-face.pt'))

init_log_file()

# job_id -> VideoJob, only while a job is actively processing this run.
# Completed jobs are always readable from disk via ANALYSIS_DIR too.
_active_jobs = {}


def _allowed_video(filename: str) -> bool:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    return ext in config.ALLOWED_VIDEO_EXTENSIONS


def _read_history(user_id=None):
    if not os.path.exists(config.HISTORY_FILE):
        return []
    try:
        with open(config.HISTORY_FILE) as f:
            history = json.load(f)
    except (json.JSONDecodeError, OSError):
        return []

    # Ensure user_id is populated for records
    modified = False
    user_map = {}
    for item in history:
        if item.get("user_id") is None:
            uname = item.get("uploaded_by")
            if uname:
                if uname not in user_map:
                    u = User.query.filter_by(username=uname).first()
                    user_map[uname] = u.id if u else None
                if user_map[uname] is not None:
                    item["user_id"] = user_map[uname]
                    modified = True

    if modified:
        try:
            with open(config.HISTORY_FILE, "w") as f:
                json.dump(history, f)
        except Exception:
            pass

    if user_id is not None:
        return [h for h in history if h.get("user_id") == user_id]

    return history


def _read_analysis(job_id: str):
    safe_id = secure_filename(job_id)
    path = os.path.join(config.ANALYSIS_DIR, f"{safe_id}.json")
    if not os.path.exists(path):
        return None
    try:
        with open(path) as f:
            data = json.load(f)
    except (json.JSONDecodeError, OSError):
        return None

    # Backfill user_id if missing
    if data and data.get("user_id") is None:
        uname = data.get("uploaded_by")
        if uname:
            u = User.query.filter_by(username=uname).first()
            if u:
                data["user_id"] = u.id
                try:
                    with open(path, "w") as f:
                        json.dump(data, f)
                except Exception:
                    pass
    return data


def _get_job_user_id(job_id: str):
    """Returns the user_id associated with a job_id, or None."""
    safe_id = secure_filename(job_id)
    # 1. Check active jobs
    job = _active_jobs.get(safe_id)
    if job and getattr(job, "user_id", None) is not None:
        return job.user_id

    # 2. Check analysis file
    analysis = _read_analysis(safe_id)
    if analysis and analysis.get("user_id") is not None:
        return analysis["user_id"]

    # 3. Check history
    for item in _read_history(user_id=None):
        if item.get("job_id") == safe_id and item.get("user_id") is not None:
            return item["user_id"]

    # 4. Check DB
    try:
        rec = UploadedVideo.query.filter_by(job_id=safe_id).first()
        if rec and rec.user_id:
            return rec.user_id
    except Exception:
        pass

    return None


def _job_belongs_to_current_user(job_id: str) -> bool:
    """Verifies whether job_id was uploaded by the currently logged-in user."""
    uid = current_user_id()
    if not uid:
        return False
    job_uid = _get_job_user_id(job_id)
    return job_uid == uid


# ---------------------------------------------------------------------------
# AUTH
# ---------------------------------------------------------------------------
@app.route("/forgot_password", methods=["GET", "POST"])
def forgot_password():

    if request.method == "POST":

        email = request.form["email"]

        user = User.query.filter_by(email=email).first()

        if not user:
            return render_template(
                "forgot_password.html",
                error="No account found with this email."
            )

        otp = str(random.randint(100000, 999999))

        otp_storage[email] = otp

        msg = MIMEText(
            f"Your Crowd Detection password reset OTP is:\n\n{otp}"
        )

        msg["Subject"] = "Password Reset OTP"
        msg["From"] = config.SENDER_EMAIL
        msg["To"] = email

        try:

            server = smtplib.SMTP_SSL(
                config.SMTP_SERVER,
                config.SMTP_PORT
            )

            server.login(
                config.SENDER_EMAIL,
                config.SENDER_PASSWORD
            )

            server.sendmail(
                config.SENDER_EMAIL,
                email,
                msg.as_string()
            )

            server.quit()

            session["reset_email"] = email

            return redirect(url_for("verify_otp"))

        except Exception as e:

            return render_template(
                "forgot_password.html",
                error=str(e)
            )

    return render_template("forgot_password.html")

@app.route("/verify_otp", methods=["GET", "POST"])
def verify_otp():

    email = session.get("reset_email")

    if not email:
        return redirect(url_for("forgot_password"))

    if request.method == "POST":

        otp = request.form["otp"]

        if otp_storage.get(email) == otp:

            session["otp_verified"] = True

            return redirect(url_for("reset_password"))

        return render_template(
            "verify_otp.html",
            error="Invalid OTP"
        )

    return render_template("verify_otp.html")

@app.route("/reset_password", methods=["GET", "POST"])
def reset_password():

    email = session.get("reset_email")

    if not email or not session.get("otp_verified"):
        return redirect(url_for("forgot_password"))

    if request.method == "POST":

        password = request.form["password"]

        confirm = request.form["confirm_password"]

        if password != confirm:

            return render_template(
                "reset_password.html",
                error="Passwords do not match"
            )

        user = User.query.filter_by(email=email).first()

        user.set_password(password)

        db.session.commit()

        otp_storage.pop(email, None)

        session.pop("reset_email", None)

        session.pop("otp_verified", None)

        return redirect(url_for("login"))

    return render_template("reset_password.html")

@app.route("/login", methods=["GET", "POST"])
def login():

    if request.method == "POST":

        username = request.form.get("username")
        password = request.form.get("password")

        user = verify_credentials(username, password)

        if user:

            session["user"] = user.username
            session["user_id"] = user.id

            return redirect(url_for("index"))

        return render_template(
            "login.html",
            error="Invalid username or password"
        )

    return render_template("login.html")

@app.route("/register", methods=["GET", "POST"])
def register():

    if request.method == "POST":

        username = request.form["username"]
        email = request.form["email"]
        phone = request.form["phone"]
        password = request.form["password"]

        if User.query.filter_by(username=username).first():
            return render_template(
                "register.html",
                error="Username already exists"
            )

        if User.query.filter_by(email=email).first():
            return render_template(
                "register.html",
                error="Email already registered"
            )

        user = User(
            # pyrefly: ignore [unexpected-keyword]
            full_name=username,
            # pyrefly: ignore [unexpected-keyword]
            username=username,
            # pyrefly: ignore [unexpected-keyword]
            email=email,
            # pyrefly: ignore [unexpected-keyword]
            phone=phone,
        )

        user.set_password(password)

        db.session.add(user)
        db.session.commit()

        return redirect(url_for("login"))

    return render_template("register.html")

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


# ---------------------------------------------------------------------------
# PAGE
# ---------------------------------------------------------------------------
@app.route("/")
@login_required
def index():
    return render_template("dashboard.html", user=current_user(), config=config)

@app.route("/emergency_contacts")
@login_required
def emergency_contacts():
    return render_template(
        "emergency_contacts.html",
        user=current_user(),
        config=config
    )
# ---------------------------------------------------------------------------
# UPLOAD + PROCESSING
# ---------------------------------------------------------------------------
@app.route("/api/upload", methods=["POST"])
@login_required
def api_upload():
    if "video" not in request.files:
        return jsonify({"error": "No file part 'video' in request."}), 400
    file = request.files["video"]
    if file.filename == "":
        return jsonify({"error": "No file selected."}), 400
    if not _allowed_video(file.filename):
        return jsonify({"error": "Unsupported format. Allowed: "
                         + ", ".join(sorted(config.ALLOWED_VIDEO_EXTENSIONS))}), 400

    job_id = new_job_id()
    safe_name = secure_filename(file.filename)
    stored_name = f"{job_id}_{safe_name}"
    save_path = os.path.join(config.UPLOAD_DIR, stored_name)
    file.save(save_path)

    threshold_mode = request.form.get("threshold_mode", "with").strip().lower()
    raw_thresh = request.form.get("threshold")

    if threshold_mode == "without" or raw_thresh in ("0", 0, "none", "null", ""):
        threshold_mode = "without"
        threshold = None
    else:
        threshold_mode = "with"
        try:
            threshold = int(raw_thresh) if raw_thresh is not None else config.CROWD_THRESHOLD
            if threshold <= 0:
                threshold_mode = "without"
                threshold = None
        except (ValueError, TypeError):
            threshold = config.CROWD_THRESHOLD

    user_id = current_user_id()
    username = current_user()

    # Save to UploadedVideo database table
    try:
        video_rec = UploadedVideo(
            job_id=job_id,
            user_id=user_id,
            original_filename=file.filename,
            stored_filename=stored_name,
            upload_path=save_path,
            status="queued",
            threshold=threshold,
            threshold_mode=threshold_mode,
        )
        db.session.add(video_rec)
        db.session.commit()
    except Exception as db_err:
        print(f"[UPLOAD] Warning: could not save UploadedVideo to DB: {db_err}")

    job = VideoJob(job_id, save_path, file.filename, _shared_model, fallback_model=_shared_fallback_model,
                    uploaded_by=username, user_id=user_id, threshold=threshold, threshold_mode=threshold_mode)
    _active_jobs[job_id] = job
    job.start()

    return jsonify({"job_id": job_id, "status": "queued"})


@app.route("/api/status/<job_id>")
@login_required
def api_status(job_id):
    if not _job_belongs_to_current_user(job_id):
        return jsonify({"error": "Unauthorized"}), 403

    job = _active_jobs.get(job_id)
    if job is not None:
        status = job.get_status()
        if status["status"] == "done":
            _active_jobs.pop(job_id, None)
        return jsonify(status)

    # Not actively running — check if it already finished in a prior request.
    analysis = _read_analysis(job_id)
    if analysis:
        return jsonify({"job_id": job_id, "status": "done", "progress": 100})
    return jsonify({"error": "Unknown job_id"}), 404


@app.route("/api/analysis/<job_id>")
@login_required
def api_analysis(job_id):
    if not _job_belongs_to_current_user(job_id):
        return jsonify({"error": "Unauthorized"}), 403

    analysis = _read_analysis(job_id)
    if not analysis:
        return jsonify({"error": "Analysis not found (still processing or invalid id)."}), 404
    return jsonify(analysis)


@app.route("/api/history")
@login_required
def api_history():
    return jsonify(_read_history(user_id=current_user_id()))


@app.route("/api/dashboard_summary")
@login_required
def api_dashboard_summary():
    history = _read_history(user_id=current_user_id())
    videos_analyzed = len(history)
    max_crowd = max((h["max_count"] for h in history), default=0)
    avg_crowd = round(sum(h["avg_count"] for h in history) / videos_analyzed, 1) if videos_analyzed else 0
    total_alerts = sum(h["total_alerts"] for h in history)
    return jsonify({
        "videos_analyzed": videos_analyzed,
        "max_crowd": max_crowd,
        "avg_crowd": avg_crowd,
        "total_alerts": total_alerts,
        "recent": history[:8],
    })


# ---------------------------------------------------------------------------
# MEDIA
# ---------------------------------------------------------------------------
@app.route("/media/processed/<path:filename>")
@login_required
def media_processed(filename):
    job_id = filename.split(".")[0].split("_")[0]
    if not _job_belongs_to_current_user(job_id):
        return jsonify({"error": "Unauthorized"}), 403
    return send_from_directory(config.PROCESSED_DIR, filename)


@app.route("/media/keyframes/<path:filename>")
@login_required
def media_keyframes(filename):
    job_id = filename.split("_")[0]
    if not _job_belongs_to_current_user(job_id):
        return jsonify({"error": "Unauthorized"}), 403
    return send_from_directory(config.KEYFRAME_DIR, filename)


@app.route("/media/heatmaps/<path:filename>")
@login_required
def media_heatmaps(filename):
    job_id = filename.split(".")[0]
    if not _job_belongs_to_current_user(job_id):
        return jsonify({"error": "Unauthorized"}), 403
    return send_from_directory(config.HEATMAP_DIR, filename)


@app.route("/media/uploads/<path:filename>")
@login_required
def media_uploads(filename):
    job_id = filename.split("_")[0]
    if not _job_belongs_to_current_user(job_id):
        return jsonify({"error": "Unauthorized"}), 403
    return send_from_directory(config.UPLOAD_DIR, filename)


@app.route("/api/preview/<job_id>")
@login_required
def api_preview(job_id):
    """Serve the latest annotated frame as JPEG for live processing preview."""
    if not _job_belongs_to_current_user(job_id):
        return "", 403
    safe_id = secure_filename(job_id)
    preview_path = os.path.join(config.PROCESSED_DIR, f"{safe_id}_preview.jpg")
    if not os.path.exists(preview_path):
        return "", 204  # No content yet
    return send_file(preview_path, mimetype="image/jpeg")


@app.route("/api/frame_data/<job_id>")
@login_required
def api_frame_data(job_id):
    if not _job_belongs_to_current_user(job_id):
        return jsonify({"error": "Unauthorized"}), 403
    safe_id = secure_filename(job_id)
    path = os.path.join(config.ANALYSIS_DIR, f"{safe_id}_frames.json")
    if not os.path.exists(path):
        return jsonify({"error": "Frame data not found."}), 404
    return send_file(path, mimetype="application/json")


# ---------------------------------------------------------------------------
# REPORTS / EXPORT
# ---------------------------------------------------------------------------
@app.route("/api/report/pdf/<job_id>")
@login_required
def report_pdf(job_id):
    if not _job_belongs_to_current_user(job_id):
        return jsonify({"error": "Unauthorized"}), 403
    safe_id = secure_filename(job_id)
    path = os.path.join(config.REPORT_DIR, f"{safe_id}.pdf")
    if not os.path.exists(path):
        return jsonify({"error": "PDF report not available for this video."}), 404
    return send_file(path, as_attachment=True,
                      download_name=f"crowd_report_{safe_id}.pdf")


@app.route("/api/report/csv/<job_id>")
@login_required
def report_csv(job_id):
    if not _job_belongs_to_current_user(job_id):
        return jsonify({"error": "Unauthorized"}), 403
    analysis = _read_analysis(job_id)
    if not analysis:
        return jsonify({"error": "Analysis not found."}), 404

    buf = io.StringIO()
    writer = csv.writer(buf)
    
    is_without = (analysis.get("summary", {}).get("threshold_mode") == "without" or analysis.get("summary", {}).get("threshold") in (None, 0, "Disabled"))
    writer.writerow(["Threshold Mode", "Without Threshold (Alerts Disabled)" if is_without else "With Threshold"])
    writer.writerow(["Safe Crowd Threshold", "Disabled" if is_without else analysis.get("summary", {}).get("threshold", config.CROWD_THRESHOLD)])
    writer.writerow([])
    writer.writerow(["Time (s)", "Count", "Status"])
    for row in analysis["timeseries"]:
        writer.writerow([row["t"], row["count"], row["status"]])
    writer.writerow([])
    writer.writerow(["Event #", "Start (s)", "End (s)", "Duration (s)",
                      "Peak Count", "Density %", "Severity", "Status"])
    for i, e in enumerate(analysis["events"], start=1):
        writer.writerow([e.get("id", e.get("event_id", i)), e.get("start_t"), e.get("end_t"), e.get("duration"),
                          e.get("peak_count"), e.get("density_pct"), e.get("severity"), e.get("status")])

    mem = io.BytesIO(buf.getvalue().encode("utf-8"))
    return send_file(mem, as_attachment=True, mimetype="text/csv",
                      download_name=f"crowd_analysis_{job_id}.csv")


@app.route("/api/report/excel/<job_id>")
@login_required
def report_excel(job_id):
    if not _job_belongs_to_current_user(job_id):
        return jsonify({"error": "Unauthorized"}), 403
    analysis = _read_analysis(job_id)
    if not analysis:
        return jsonify({"error": "Analysis not found."}), 404
    try:
        from openpyxl import Workbook
    except ImportError:
        return jsonify({"error": "openpyxl is not installed on the server "
                         "(pip install openpyxl)."}), 500

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Summary"
    ws1.append(["Video", analysis["original_filename"]])
    is_without = (analysis.get("summary", {}).get("threshold_mode") == "without" or analysis.get("summary", {}).get("threshold") in (None, 0, "Disabled"))
    ws1.append(["Threshold Mode", "Without Threshold (Alerts Disabled)" if is_without else "With Threshold"])
    ws1.append(["Safe Crowd Threshold", "Disabled" if is_without else analysis.get("summary", {}).get("threshold", config.CROWD_THRESHOLD)])
    ws1.append(["Duration (s)", analysis["video"]["duration_sec"]])
    ws1.append(["Resolution", analysis["video"]["resolution"]])
    ws1.append(["Average Crowd", analysis["summary"]["avg_count"]])
    ws1.append(["Maximum Crowd", analysis["summary"]["max_count"]])
    ws1.append(["Peak Time (s)", analysis["summary"]["peak_time_sec"]])
    ws1.append(["Total Alerts", analysis["summary"]["total_alerts"]])
    ws1.append(["Overall Risk", analysis["summary"]["overall_risk"]])

    ws2 = wb.create_sheet("Timeseries")
    ws2.append(["Time (s)", "Count", "Status"])
    for row in analysis["timeseries"]:
        ws2.append([row["t"], row["count"], row["status"]])

    ws3 = wb.create_sheet("Events")
    ws3.append(["Event #", "Start (s)", "End (s)", "Duration (s)",
                "Peak Count", "Density %", "Severity", "Status"])
    for i, e in enumerate(analysis["events"], start=1):
        ws3.append([e.get("id", e.get("event_id", i)), e.get("start_t"), e.get("end_t"), e.get("duration"),
                    e.get("peak_count"), e.get("density_pct"), e.get("severity"), e.get("status")])

    mem = io.BytesIO()
    wb.save(mem)
    mem.seek(0)
    return send_file(mem, as_attachment=True,
                      mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                      download_name=f"crowd_analysis_{job_id}.xlsx")


# ---------------------------------------------------------------------------
# EMERGENCY CONTACTS API
# ---------------------------------------------------------------------------

@app.route("/api/emergency_contacts", methods=["GET", "POST"])
@login_required
def api_emergency_contacts():
    if request.method == "GET":
        contacts = EmergencyContact.query.all()
        return jsonify([c.to_dict() for c in contacts])
    # POST - create new contact
    data = request.get_json() or {}
    contact = EmergencyContact(
        # pyrefly: ignore [unexpected-keyword]
        name=data.get("name"),
        # pyrefly: ignore [unexpected-keyword]
        phone=data.get("phone"),
        # pyrefly: ignore [unexpected-keyword]
        email=data.get("email"),
        # pyrefly: ignore [unexpected-keyword]
        role=data.get("role"),
        # pyrefly: ignore [unexpected-keyword]
        priority=data.get("priority", 1),
        # pyrefly: ignore [unexpected-keyword]
        camera_zone=data.get("camera_zone"),
    )
    db.session.add(contact)
    db.session.commit()
    return jsonify(contact.to_dict()), 201

@app.route("/api/emergency_contacts/<int:contact_id>", methods=["PUT", "DELETE"])
@login_required
def api_emergency_contact_detail(contact_id):
    contact = EmergencyContact.query.get_or_404(contact_id)
    if request.method == "PUT":
        data = request.get_json() or {}
        for field in ["name", "phone", "email", "role", "priority", "camera_zone"]:
            if field in data:
                setattr(contact, field, data[field])
        db.session.commit()
        return jsonify(contact.to_dict())
    # DELETE
    db.session.delete(contact)
    db.session.commit()
    return "", 204

# ---------------------------------------------------------------------------
# ALERT SETTINGS API
# ---------------------------------------------------------------------------

@app.route("/api/alert_settings", methods=["GET", "POST"])
@login_required
def api_alert_settings():
    if request.method == "GET":
        setting = AlertSetting.query.first()
        if not setting:
            setting = AlertSetting()
            db.session.add(setting)
            db.session.commit()
        return jsonify(setting.to_dict())
    # POST - update/create settings
    data = request.get_json() or {}
    setting = AlertSetting.query.first()
    if not setting:
        setting = AlertSetting()
        db.session.add(setting)
    for field in ["threshold", "notify_sms", "notify_email", "notify_whatsapp", "cooldown_minutes"]:
        if field in data:
            setattr(setting, field, data[field])
    db.session.commit()
    return jsonify(setting.to_dict())

# ---------------------------------------------------------------------------
# ALERT HISTORY API
# ---------------------------------------------------------------------------

@app.route("/api/alert_history", methods=["GET"])
@login_required
def api_alert_history():
    history = AlertHistory.query.order_by(AlertHistory.timestamp.desc()).all()
    return jsonify([h.to_dict() for h in history])

# ---------------------------------------------------------------------------
# TEST ALERT ENDPOINT
# ---------------------------------------------------------------------------

@app.route("/api/test_alert", methods=["POST"])
@login_required
def api_test_alert():
    # Simple test payload
    details = {
        "Location": "Test Zone",
        "Camera Name": "Test Camera",
        "Time": datetime.utcnow().isoformat(),
        "Crowd Count": 99,
        "Threshold": 50,
    }
    alerts.send_alert(camera_name="Test Camera", severity="WARNING", details=details)
    return jsonify({"status": "test alert sent"})

# ---------------------------------------------------------------------------
# PUBLIC CROWD SAFETY ADVISORY & QR CODE
# ---------------------------------------------------------------------------

def get_lan_ip():
    """Detect the local machine's LAN IP address so mobile phones on the Wi-Fi network can connect."""
    if getattr(config, "PUBLIC_BASE_URL", None):
        return config.PUBLIC_BASE_URL.rstrip("/")
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(('8.8.8.8', 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        try:
            return socket.gethostbyname(socket.gethostname())
        except Exception:
            return "127.0.0.1"


def get_public_base_url():
    """
    Returns the root base URL (e.g. http://10.146.107.198:5000) accessible
    from mobile devices on the local network or configured public URL.
    """
    if getattr(config, "PUBLIC_BASE_URL", None):
        return config.PUBLIC_BASE_URL.rstrip("/")

    # If the user accessed the dashboard through an IP other than localhost/127.0.0.1,
    # preserve that host (including port)
    if request:
        try:
            req_host = request.host
            hostname = req_host.split(":")[0]
            if hostname not in ("localhost", "127.0.0.1", "0.0.0.0"):
                scheme = request.scheme or "http"
                return f"{scheme}://{req_host}"
        except Exception:
            pass

    lan_ip = get_lan_ip()
    port = getattr(config, "DASHBOARD_PORT", 5000)
    return f"http://{lan_ip}:{port}"


def get_crowd_status_data(job_id=None, user_id=None):
    """
    Retrieves latest crowd safety data for an active or completed job,
    or the most recent analysis for user_id if job_id is omitted.
    """
    active_job = None
    analysis = None
    target_job_id = None

    if job_id:
        safe_id = secure_filename(job_id)
        # Check active jobs (exact or prefix match)
        for jid, j in _active_jobs.items():
            if jid == safe_id or jid.startswith(safe_id):
                active_job = j
                target_job_id = jid
                break
        if not active_job:
            # Check analysis directory
            analysis = _read_analysis(safe_id)
            if not analysis and os.path.exists(config.ANALYSIS_DIR):
                for fn in os.listdir(config.ANALYSIS_DIR):
                    if fn.endswith(".json") and not fn.endswith("_frames.json") and fn.startswith(safe_id):
                        target_job_id = fn[:-5]
                        analysis = _read_analysis(target_job_id)
                        break
            else:
                target_job_id = safe_id
    elif user_id is not None:
        # Check active jobs for this user
        for jid, j in reversed(list(_active_jobs.items())):
            if getattr(j, "user_id", None) == user_id:
                target_job_id = jid
                active_job = j
                break
        if not active_job:
            user_history = _read_history(user_id=user_id)
            if user_history:
                target_job_id = user_history[0]["job_id"]
                analysis = _read_analysis(target_job_id)
    else:
        # Neither job_id nor user_id provided
        pass

    # Determine if without threshold
    is_without_threshold = False
    if active_job:
        is_without_threshold = (getattr(active_job, "threshold_mode", "with") == "without")
    elif analysis:
        summary = analysis.get("summary", {})
        is_without_threshold = (summary.get("threshold_mode") == "without" or summary.get("threshold") in (None, 0, "Disabled"))

    # Extract crowd figures
    if active_job:
        status_info = active_job.get_status()
        current_count = status_info.get("current_count", 0)
        safe_capacity = getattr(active_job, "threshold", None)
        is_live = True
        video_filename = getattr(active_job, "original_filename", "")
    elif analysis:
        summary = analysis.get("summary", {})
        current_count = summary.get("max_count", 0)
        safe_capacity = summary.get("threshold", None)
        is_live = False
        video_filename = analysis.get('original_filename', '')
    else:
        # Baseline fallback when no video exists for this user
        current_count = 0
        safe_capacity = config.CROWD_THRESHOLD
        is_live = False
        video_filename = "No video uploaded yet"

    current_count = int(current_count)

    if is_without_threshold:
        status = "MONITORING"
        safe_capacity_display = "Disabled"
        occupancy_pct = "N/A"
        message = "Crowd monitoring active. Threshold-based capacity alerts are disabled."
    else:
        safe_cap_val = max(1, int(safe_capacity or config.CROWD_THRESHOLD))
        safe_capacity_display = safe_cap_val
        occupancy_pct = round((current_count / safe_cap_val) * 100, 1)

        # Classify status (SAFE / WARNING / CRITICAL)
        if not active_job and not analysis:
            status = "SAFE"
            message = "No video analysis active. Upload a video to monitor crowd safety."
        elif current_count >= safe_cap_val:
            status = "CRITICAL"
            message = "Please do not enter. The crowd level has exceeded the safe capacity."
        elif current_count >= round(safe_cap_val * 0.7):
            status = "WARNING"
            message = "Caution advised. High crowd density detected. Please follow staff instructions."
        else:
            status = "SAFE"
            message = "Normal crowd density. Entry permitted. Please proceed safely."

    # Format timestamp: e.g. 8:22:56 PM
    now = datetime.now()
    last_updated = now.strftime("%I:%M:%S %p").lstrip("0")

    return {
        "job_id": target_job_id or job_id or "",
        "user_id": user_id,
        "threshold_mode": "without" if is_without_threshold else "with",
        "status": status,
        "current_people": current_count,
        "current_count": current_count,
        "safe_capacity": safe_capacity_display,
        "occupancy_percentage": occupancy_pct,
        "occupancy_pct": occupancy_pct,
        "last_updated": last_updated,
        "message": message,
        "is_active": is_live,
        "is_live": is_live,
        "video_filename": video_filename,
    }


@app.route("/public")
@app.route("/public/<job_id>")
@app.route("/public/user/<int:user_id>")
def public_crowd_status(job_id=None, user_id=None):
    """
    Public Crowd Safety Status webpage accessible via mobile phone QR scan.
    No login required.
    """
    if user_id is None and not job_id and session.get("user_id"):
        user_id = session.get("user_id")
    data = get_crowd_status_data(job_id=job_id, user_id=user_id)
    initial_cap = data["safe_capacity"] if data.get("threshold_mode") != "without" else "Disabled"
    initial_occ = f"{data['occupancy_percentage']}%" if data.get("threshold_mode") != "without" else "N/A"
    return render_template(
        "public_status.html",
        job_id=data.get("job_id", ""),
        user_id=user_id or data.get("user_id", ""),
        initial_status=data["status"],
        initial_count=data["current_people"],
        initial_capacity=initial_cap,
        initial_occupancy=initial_occ,
        initial_time=data["last_updated"],
        video_filename=data.get("video_filename", ""),
    )


@app.route("/api/public_status")
@app.route("/api/public_status/<job_id>")
@app.route("/api/public_status/user/<int:user_id>")
def api_public_status(job_id=None, user_id=None):
    """JSON API endpoint for the public page to auto-refresh crowd metrics."""
    if user_id is None and not job_id and session.get("user_id"):
        user_id = session.get("user_id")
    data = get_crowd_status_data(job_id=job_id, user_id=user_id)
    return jsonify(data)


@app.route("/api/qr_code")
@app.route("/api/qr_code/<job_id>")
@app.route("/api/qr_code/user/<int:user_id>")
def api_qr_code(job_id=None, user_id=None):
    """Generates a dynamic QR code PNG pointing to the public crowd status URL."""
    base_url = get_public_base_url()
    if job_id:
        target_url = f"{base_url}/public/{job_id}"
    elif user_id:
        target_url = f"{base_url}/public/user/{user_id}"
    elif session.get("user_id"):
        uid = session.get("user_id")
        user_data = get_crowd_status_data(user_id=uid)
        if user_data.get("job_id"):
            target_url = f"{base_url}/public/{user_data['job_id']}"
        else:
            target_url = f"{base_url}/public/user/{uid}"
    else:
        target_url = f"{base_url}/public"

    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=10,
        border=3,
    )
    qr.add_data(target_url)
    qr.make(fit=True)
    img = qr.make_image(fill_color="#0f172a", back_color="#ffffff")

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return send_file(buf, mimetype="image/png")


@app.route("/api/public_info")
@app.route("/api/public_info/<job_id>")
def api_public_info(job_id=None):
    """Returns URL and metadata for sharing / displaying QR codes on the dashboard."""
    base_url = get_public_base_url()
    uid = current_user_id()

    if job_id:
        if uid and not _job_belongs_to_current_user(job_id):
            return jsonify({"error": "Unauthorized"}), 403
        public_url = f"{base_url}/public/{job_id}"
        qr_url = f"/api/qr_code/{job_id}"
        info_data = get_crowd_status_data(job_id=job_id, user_id=uid)
    else:
        info_data = get_crowd_status_data(user_id=uid)
        target_job_id = info_data.get("job_id")
        if target_job_id:
            public_url = f"{base_url}/public/{target_job_id}"
            qr_url = f"/api/qr_code/{target_job_id}"
        elif uid:
            public_url = f"{base_url}/public/user/{uid}"
            qr_url = f"/api/qr_code/user/{uid}"
        else:
            public_url = f"{base_url}/public"
            qr_url = "/api/qr_code"

    return jsonify({
        "public_url": public_url,
        "qr_code_url": qr_url,
        "job_id": info_data.get("job_id", ""),
        "user_id": uid,
        "lan_ip": get_lan_ip(),
        "video_filename": info_data.get('video_filename', ''),
    })


if __name__ == "__main__":
    app.run(host=config.DASHBOARD_HOST, port=config.DASHBOARD_PORT,
             debug=False, threaded=True)
